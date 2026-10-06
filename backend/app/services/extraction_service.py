import re
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from app.core.exceptions import PrepXException


class ResumeExtractorService:
    """Extracts raw plain text cleanly from PDF and DOCX resume files."""

    @staticmethod
    def extract_text(file_path_str: str, file_type: str) -> str:
        path = Path(file_path_str)
        if not path.exists():
            raise PrepXException(status_code=404, detail="Resume file not found on disk", code="FILE_NOT_FOUND")

        ft = file_type.lower().replace(".", "")
        if ft == "pdf":
            return ResumeExtractorService._extract_pdf(path)
        elif ft in {"docx", "doc"}:
            return ResumeExtractorService._extract_docx(path)
        else:
            raise PrepXException(status_code=400, detail=f"Unsupported file type: {file_type}", code="INVALID_FILE_TYPE")

    @staticmethod
    def _extract_pdf(path: Path) -> str:
        try:
            reader = PdfReader(str(path))
            extracted_pages = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    extracted_pages.append(text)
            full_text = "\n".join(extracted_pages)
            return ResumeExtractorService._clean_text(full_text)
        except Exception as e:
            raise PrepXException(
                status_code=422,
                detail=f"Failed to extract text from PDF document: {str(e)}",
                code="PDF_EXTRACTION_FAILED",
            )

    @staticmethod
    def _extract_docx(path: Path) -> str:
        try:
            doc = Document(str(path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        paragraphs.append(" | ".join(row_texts))
            full_text = "\n".join(paragraphs)
            return ResumeExtractorService._clean_text(full_text)
        except Exception as e:
            raise PrepXException(
                status_code=422,
                detail=f"Failed to extract text from DOCX document: {str(e)}",
                code="DOCX_EXTRACTION_FAILED",
            )

    @staticmethod
    def _clean_text(text: str) -> str:
        # Normalize multiple spaces, tabs, and duplicate blank lines
        text = re.sub(r"[ \t]+", " ", text)
        text = re.sub(r"\n\s*\n+", "\n\n", text)
        return text.strip()
